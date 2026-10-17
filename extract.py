"""Stage 4: turn the downloaded course files into clean, searchable text.

    python extract.py            # extract new or changed files, rebuild indexes and index.html
    python extract.py --force    # re-extract everything

Reads data/manifest.json (written by scrape.py) and never contacts Moodle.
Writes, for course <course>:
  data/text/<course>/<NN-section>/<file>.md   one per source, with page / slide / ¶ locations
  data/text/<course>/<NN-section>/_index.md   per-section index
  data/text/<course>/_index.md                course index
  data/text/<course>/passages.jsonl           search passages for Stage 5
  data/text/<course>/extract_state.json       cache: what was extracted from which file version
  downloads/<course>/index.html               offline reading view of the whole course
Formats and rules: .claude/skills/text-extraction/SKILL.md
"""

import base64
import hashlib
import html
import json
import os
import re
import subprocess
import sys
import unicodedata
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import quote

from common import (DATA_DIR, MANIFEST_FILE, ROOT, ConfigError, base_parser,
                    course_dir, log, run, section_dir, write_json_atomic,
                    write_text)

# Bump when extraction output changes, so every file is re-extracted once.
EXTRACTOR_VERSION = 3
TEXT_DIR = DATA_DIR / "text"
CONVERTED_DIR = DATA_DIR / "converted"
SCANNED_PAGE_CHARS = 25         # a PDF page with less text than this, plus an image, counts as scanned
PASSAGE_CHARS = 1200            # Word paragraphs are grouped into search passages up to about this size
PPT_TIMEOUT_S = 120
# Tests point this at a missing program to simulate "PowerPoint not available".
POWERSHELL = os.environ.get("MOODLE_STUDY_POWERSHELL", "powershell")

KINDS = {".pdf": "pdf", ".pptx": "pptx", ".ppt": "ppt", ".docx": "docx", ".md": "md",
         ".html": "html", ".htm": "html", ".txt": "txt"}
MEDIA = {".mp4": "video", ".mov": "video", ".avi": "video", ".mkv": "video", ".webm": "video", ".wmv": "video",
         ".mp3": "audio", ".m4a": "audio", ".wav": "audio",
         ".png": "image", ".jpg": "image", ".jpeg": "image", ".gif": "image"}
OLD_OFFICE = {".doc": ".docx", ".xls": ".xlsx", ".pps": ".pptx"}
UNITS = {"pdf": "page", "pptx": "slide", "ppt": "slide", "docx": "paragraph"}
KIND_LABEL = {"pdf": "PDF", "pptx": "PowerPoint", "ppt": "PowerPoint (.ppt)", "docx": "Word", "md": "Instructions",
              "html": "Web page", "txt": "Text", "video": "Video", "audio": "Audio", "image": "Image",
              "old_office": "Old Office file", "other": "Other"}

# Flags worth the user's attention, with the advice shown in indexes and index.html.
FLAG_TEXT = {
    "needs_ocr": "Scanned: little or no text layer. Run ocrmypdf on it (needs Tesseract) to make it searchable.",
    "mostly_images": "Most slides are images only, so their content isn't searchable.",
    "needs_conversion": "Old Office format that couldn't be converted automatically. Open it and Save As the "
                        "modern format, then rerun extract.py.",
    "unsupported": "File type not supported for text extraction.",
    "error": "Couldn't be read.",
}
INFO_FLAGS = {"no_text": "Media file: no text to extract."}
RETRY_FLAGS = {"error", "needs_conversion"}  # re-attempted on every run


class ExtractError(Exception):
    pass


# --- text helpers ---------------------------------------------------------------

_SPACES = re.compile(r"[\u00a0\u2000-\u200a\u202f\u205f\u3000\t]")
_CTRL = re.compile(r"[\x00-\x08\x0c\x0e-\x1f\x7f\u200b\ufeff]")  # keeps ZWJ/ZWNJ (needed for Arabic/Persian)


def clean_inline(s: str) -> str:
    s = unicodedata.normalize("NFC", s or "")
    s = s.replace("\r\n", "\n").replace("\r", "\n").replace("\v", "\n")
    return _CTRL.sub("", _SPACES.sub(" ", s))


def clean(s: str) -> str:
    lines = (re.sub(r" {2,}", " ", line).strip() for line in clean_inline(s).split("\n"))
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def plain(md: str) -> str:
    return md.replace("**", "")


def fmt_runs(runs, bold_of) -> str:
    """Join text runs, wrapping bold runs in **...** (flashcard candidates for Stage 5)."""
    out = []
    for r in runs:
        t = clean_inline(r.text).replace("\n", " ")
        if not t:
            continue
        if bold_of(r) and t.strip() and "*" not in t:
            core = t.strip()
            out.append(t[:len(t) - len(t.lstrip())] + f"**{core}**" + t[len(t.rstrip()):])
        else:
            out.append(t)
    return re.sub(r" {2,}", " ", "".join(out).replace("****", "")).strip()


def md_table(rows: list[list[str]]) -> str:
    rows = [[c.replace("|", "\\|").replace("\n", " / ") for c in r] for r in rows if any(c.strip() for c in r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    return "\n".join(lines + ["| " + " | ".join(r) + " |" for r in rows[1:]])


def block(loc: str, n: int, md: str, heading: str = "", level: int = 0, text: str | None = None) -> dict:
    return {"loc": loc, "n": n, "heading": heading, "level": level, "md": md,
            "text": plain(md) if text is None else text}


def words(blocks: list[dict]) -> int:
    return sum(len(b["text"].split()) for b in blocks)


# --- extractors: each returns {"blocks", "flags", "outline", "note"} -------------

def extract_pdf(path: Path) -> dict:
    import pymupdf
    try:
        doc = pymupdf.open(str(path))
    except Exception as e:
        raise ExtractError(f"not a readable PDF ({e})")
    if doc.needs_pass:
        raise ExtractError("the PDF is password-protected")
    blocks, scanned = [], 0
    for i, page in enumerate(doc, 1):
        raw = re.sub(r"(\w)-\n([a-z])", r"\1\2", page.get_text("text", sort=True))  # re-join hyphenated words
        text = clean(raw)
        if len(text) < SCANNED_PAGE_CHARS and page.get_images():
            scanned += 1
        blocks.append(block(f"page {i}", i, text or "*(No text layer on this page.)*", text=text))
    flags = ["needs_ocr"] if doc.page_count and scanned * 2 >= doc.page_count else []
    outline = [clean(t[1]) for t in doc.get_toc()[:30]]
    note = f"{scanned} of {doc.page_count} pages look scanned." if scanned else ""
    return {"blocks": blocks, "flags": flags, "outline": outline, "note": note}


_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_DGM_RELIDS = "{http://schemas.openxmlformats.org/drawingml/2006/diagram}relIds"
_R_DM = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}dm"


def smartart_paragraphs(shape, slide) -> list[str]:
    """Text of a SmartArt diagram. It lives in a separate diagram-data part that python-pptx doesn't read."""
    from lxml import etree
    rel = shape._element.find(f".//{_DGM_RELIDS}")
    if rel is None or not rel.get(_R_DM):
        return []
    try:
        root = etree.fromstring(slide.part.related_part(rel.get(_R_DM)).blob)
    except (KeyError, etree.XMLSyntaxError):
        return []
    paras = (clean("".join(t.text or "" for t in p.iter(f"{_A}t"))) for p in root.iter(f"{_A}p"))
    return [t.replace("\n", " ") for t in paras if t]


def extract_pptx(path: Path) -> dict:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
    BOILERPLATE = {PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER}
    try:
        prs = Presentation(str(path))
    except Exception as e:
        raise ExtractError(f"not a readable .pptx ({e})")

    def leaves(shapes):
        """Shapes in reading order (top to bottom, then left to right), descending into groups."""
        for shape in sorted(shapes, key=lambda s: (s.top or 0, s.left or 0)):
            try:
                is_group = shape.shape_type == MSO_SHAPE_TYPE.GROUP
            except NotImplementedError:
                is_group = False
            if is_group:
                yield from leaves(shape.shapes)
            else:
                yield shape

    blocks, outline, textless = [], [], 0
    for n, slide in enumerate(prs.slides, 1):
        title_shape = slide.shapes.title
        title = clean(title_shape.text_frame.text) if title_shape is not None and title_shape.has_text_frame else ""
        title = title.replace("\n", " ")
        parts = []
        for shape in leaves(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if shape.is_placeholder and shape.placeholder_format.type in BOILERPLATE:
                continue  # slide number, date, footer
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = fmt_runs(p.runs, lambda r: r.font.bold) or clean(p.text).replace("\n", " ")
                    if t:
                        parts.append("  " * (p.level or 0) + "- " + t)
            elif getattr(shape, "has_table", False) and shape.has_table:
                rows = [["" if c.is_spanned else clean(c.text) for c in row.cells] for row in shape.table.rows]
                t = md_table(rows)
                if t:
                    parts.append("\n" + t + "\n")
            else:
                parts += ["- " + t for t in smartart_paragraphs(shape, slide)]
        body ="\n".join(parts).strip()
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = clean(slide.notes_slide.notes_text_frame.text)
        if not (title or body or notes):
            textless += 1
            md = "*(No text on this slide: images or diagrams only.)*"
        else:
            md = body + (f"\n\n### Notes\n\n{notes}" if notes else "")
        blocks.append(block(f"slide {n}", n, md.strip(), heading=title,
                            text=" ".join(x for x in (title, plain(body), notes) if x)))
        outline.append(title or f"(slide {n})")
    flags = ["mostly_images"] if blocks and textless * 2 > len(blocks) else []
    note = f"{textless} of {len(blocks)} slides have no text." if textless else ""
    return {"blocks": blocks, "flags": flags, "outline": outline, "note": note}


def extract_docx(path: Path) -> dict:
    import docx
    from docx.table import Table
    try:
        d = docx.Document(str(path))
    except Exception as e:
        raise ExtractError(f"not a readable .docx ({e})")
    blocks, outline, heading, n = [], [], "", 0
    for el in d.iter_inner_content():
        if isinstance(el, Table):
            rows = []
            for row in el.rows:
                cells, prev = [], None
                for c in row.cells:
                    if c._tc is prev:  # horizontally merged cell repeats
                        continue
                    prev = c._tc
                    cells.append(clean(c.text))
                rows.append(cells)
            t = md_table(rows)
            if t:
                n += 1
                blocks.append(block(f"¶{n}", n, t, heading=heading))
            continue
        text = clean(el.text)
        if not text:
            continue
        n += 1
        style = (el.style.name if el.style is not None else "") or ""
        m = re.match(r"Heading (\d)", style)
        level = 1 if style == "Title" else int(m.group(1)) if m else 0
        if level:
            heading = text.replace("\n", " ")
            outline.append("  " * (level - 1) + heading)
            blocks.append(block(f"¶{n}", n, heading, heading=heading, level=level))
            continue
        md = fmt_runs(el.runs, lambda r: r.bold)
        if plain(md).replace(" ", "") != text.replace("\n", "").replace(" ", ""):
            md = text  # runs miss some text (e.g. inside hyperlinks): keep the full text, lose bold
        pPr = el._p.pPr
        if "List" in style or (pPr is not None and pPr.numPr is not None):
            md = "- " + md
        blocks.append(block(f"¶{n}", n, md, heading=heading))
    return {"blocks": blocks, "flags": [], "outline": outline, "note": ""}


def extract_md(path: Path) -> dict:
    text = clean(path.read_text(encoding="utf-8", errors="replace"))
    body = re.sub(r"\A# [^\n]*\n*", "", text)  # the header repeats the title
    return {"blocks": [block("text", 1, body)] if body else [], "flags": [], "outline": [], "note": ""}


def extract_txt(path: Path) -> dict:
    text = clean(path.read_text(encoding="utf-8", errors="replace"))
    return {"blocks": [block("text", 1, text)] if text else [], "flags": [], "outline": [], "note": ""}


class _HTMLText(HTMLParser):
    BLOCK = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "table"}

    def __init__(self):
        super().__init__()
        self.out, self.skip = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self.skip += 1
        elif tag in self.BLOCK:
            self.out.append("\n")
        if tag == "li":
            self.out.append("- ")

    def handle_endtag(self, tag):
        if tag in ("script", "style"):
            self.skip = max(0, self.skip - 1)
        elif tag in self.BLOCK:
            self.out.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.out.append(data)


def extract_html(path: Path) -> dict:
    parser = _HTMLText()
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    text = clean(re.sub(r"[ \t]*\n[ \t]*", "\n", "".join(parser.out)))
    return {"blocks": [block("text", 1, text)] if text else [], "flags": [], "outline": [], "note": ""}


# --- .ppt -> .pptx through the installed PowerPoint ------------------------------

# Runs hidden (no window). Quits PowerPoint afterwards only if it wasn't already open with
# your own presentations, so it never closes your work.
_PS_CONVERT = r"""
$ErrorActionPreference = 'Stop'
$app = New-Object -ComObject PowerPoint.Application
$before = $app.Presentations.Count
try {
  $p = $app.Presentations.Open($env:MS_SRC, -1, 0, 0)
  try { $p.SaveAs($env:MS_DST, 24) } finally { $p.Close() }
} finally {
  if ($before -eq 0 -and $app.Presentations.Count -eq 0) { $app.Quit() }
  [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($app)
}
"""


def convert_ppt(src: Path, dst: Path) -> None:
    if dst.exists() and dst.stat().st_mtime_ns >= src.stat().st_mtime_ns:
        return
    if sys.platform != "win32":
        raise ExtractError("converting .ppt needs PowerPoint on Windows")
    dst.parent.mkdir(parents=True, exist_ok=True)
    tmp = dst.with_name(dst.stem + ".part.pptx")
    env = {**os.environ, "MS_SRC": str(src.resolve()), "MS_DST": str(tmp.resolve())}
    script = base64.b64encode(_PS_CONVERT.encode("utf-16-le")).decode()
    log.info("  converting with PowerPoint (in the background)...")
    try:
        proc = subprocess.run([POWERSHELL, "-NoProfile", "-NonInteractive", "-EncodedCommand", script],
                              env=env, capture_output=True, text=True, timeout=PPT_TIMEOUT_S)
    except FileNotFoundError:
        raise ExtractError("PowerShell wasn't found, so PowerPoint couldn't be used")
    except subprocess.TimeoutExpired:
        raise ExtractError(f"PowerPoint didn't finish within {PPT_TIMEOUT_S} s")
    if proc.returncode != 0 or not tmp.exists():
        lines = [ln.strip() for ln in (proc.stderr or proc.stdout or "").splitlines() if ln.strip()]
        reason = lines[0] if lines else f"exit code {proc.returncode}"
        tmp.unlink(missing_ok=True)
        raise ExtractError(f"PowerPoint couldn't convert it ({reason[:200]})")
    tmp.replace(dst)


# --- sources ----------------------------------------------------------------------

def rel(path: Path) -> str:
    return path.relative_to(ROOT).as_posix()


def collect_sources(manifest: dict, cdir: Path) -> tuple[list[dict], list[str]]:
    """Every downloaded file, with the manifest info it belongs to, in manifest order."""
    sources, seen, missing = [], set(), []
    for item in manifest["items"]:
        if item["status"] != "downloaded" or item["type"] == "url":
            continue
        for p in item.get("files") or ([item["local_path"]] if item["local_path"] else []):
            path = ROOT / p
            if p in seen or path.is_dir():
                continue
            seen.add(p)
            if not path.exists():
                missing.append(p)
                continue
            if item["module"] == "assign":
                title = f"{item['title']} (instructions)" if path.suffix == ".md" else f"{item['title']}: {path.name}"
            elif item["module"] == "folder":
                title = f"{item['title']}: {path.relative_to(ROOT / item['local_path']).as_posix()}"
            else:
                title = item["title"]
            sources.append({
                "key": p, "path": path, "item": item, "title": title,
                "section_number": item["section_number"], "section_name": item["section_name"],
                "section_dir": section_dir(manifest, item).name,
                "rel_course": path.relative_to(cdir).as_posix(),
            })
    return sources, missing


def kind_of(path: Path) -> str:
    ext = path.suffix.lower()
    return KINDS.get(ext) or MEDIA.get(ext) or ("old_office" if ext in OLD_OFFICE else "other")


def sha1(path: Path) -> str:
    h = hashlib.sha1()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def output_path(text_root: Path, src: dict, kind: str) -> Path | None:
    if kind not in ("pdf", "pptx", "ppt", "docx", "md", "html", "txt"):
        return None
    out = text_root / src["rel_course"]
    return out if kind == "md" else out.with_name(out.name + ".md")


def extract_source(src: dict, kind: str, converted_root: Path) -> dict:
    path = src["path"]
    result = {"blocks": [], "flags": [], "outline": [], "note": "", "converted": ""}
    try:
        if kind == "ppt":
            dst = (converted_root / src["rel_course"]).with_suffix(".pptx")
            convert_ppt(path, dst)
            result.update(extract_pptx(dst), converted=rel(dst))
        elif kind in ("video", "audio", "image"):
            result.update(flags=["no_text"], note=INFO_FLAGS["no_text"] if kind != "image" else
                          "Image: no text layer (OCR could read text in it).")
        elif kind == "old_office":
            new = OLD_OFFICE[path.suffix.lower()]
            result.update(flags=["needs_conversion"], note=f"Open it and Save As {new}, then rerun extract.py.")
        elif kind == "other":
            result.update(flags=["unsupported"], note=f"{path.suffix or 'No extension'}: not a supported format.")
        else:
            result.update({"pdf": extract_pdf, "pptx": extract_pptx, "docx": extract_docx, "md": extract_md,
                           "html": extract_html, "txt": extract_txt}[kind](path))
    except ExtractError as e:
        result.update(blocks=[], flags=["needs_conversion" if kind == "ppt" else "error"], note=str(e))
    except Exception as e:  # a broken file must never stop the run
        result.update(blocks=[], flags=["error"], note=f"{type(e).__name__}: {e}")
    return result


# --- Markdown output ---------------------------------------------------------------

def render_md(src: dict, entry: dict) -> str:
    item, kind = src["item"], entry["kind"]
    unit = UNITS.get(kind, "")
    label = KIND_LABEL[kind] + (" (converted to .pptx)" if entry.get("converted") else "")
    rows = [("Section", f"{src['section_number']:02d} {src['section_name']}"),
            ("Source", f"`{src['key']}`"), ("Moodle", item["url"]), ("Type", label)]
    if unit:
        rows.append((unit.capitalize() + "s", str(len(entry["blocks"]))))
    rows.append(("Words", f"{entry['words']:,}"))
    if entry["flags"]:
        rows.append(("Flags", ", ".join(entry["flags"])))
    out = [f"# {src['title']}", "", "| | |", "|---|---|"]
    out += [f"| {k} | {v.replace('|', chr(92) + '|')} |" for k, v in rows]
    out.append("")
    for f in entry["flags"]:
        out += [f"> **{f}:** {FLAG_TEXT.get(f) or INFO_FLAGS.get(f, '')} {entry['note']}".rstrip(), ""]
    if kind == "docx":
        out += ["> Word files have no fixed pages: `[¶N]` marks paragraph N, counted from the top.", ""]
    for b in entry["blocks"]:
        if kind == "pdf":
            out += [f"## Page {b['n']}", "", b["md"], ""]
        elif kind in ("pptx", "ppt"):
            out += [f"## Slide {b['n']}: {b['heading'] or '(untitled)'}", "", b["md"], ""]
        elif kind == "docx":
            if b["level"]:
                out += ["#" * min(b["level"] + 1, 6) + " " + b["md"] + f" [¶{b['n']}]", ""]
            elif b["md"].startswith("|"):  # a table: the marker can't go on its last row
                out += [f"[¶{b['n']}]", "", b["md"], ""]
            else:
                out += [f"{b['md']} [¶{b['n']}]", ""]
        else:
            out += [b["md"], ""]
    return re.sub(r"\n{3,}", "\n\n", "\n".join(out)).rstrip() + "\n"


def write_if_changed(path: Path, text: str) -> bool:
    if path.exists() and path.read_text(encoding="utf-8") == text:
        return False
    write_text(path, text)
    return True


def md_link(target: Path, from_dir: Path) -> str:
    return quote(os.path.relpath(target, from_dir).replace(os.sep, "/"))


def write_indexes(manifest: dict, text_root: Path, sources: list[dict], state: dict) -> None:
    by_section: dict[int, list[dict]] = {}
    for src in sources:
        by_section.setdefault(src["section_number"], []).append(src)
    course_lines = [f"# {manifest['course_name']}", "",
                    f"Extracted text of every downloaded file. Course id {manifest['course_id']}.", "",
                    "| Section | Files | Words |", "|---|---|---|"]
    attention = []
    for sec in manifest["sections"]:
        srcs = by_section.get(sec["section_number"], [])
        if not srcs:
            continue
        sec_dir = text_root / srcs[0]["section_dir"]
        lines = [f"# {sec['section_number']:02d} {sec['section_name']}", "",
                 "| File | Type | Size | Words | Flags |", "|---|---|---|---|---|"]
        outlines = []
        for src in srcs:
            e = state[src["key"]]
            name = src["title"].replace("|", "\\|")
            link = f"[{name}]({md_link(ROOT / e['output'], sec_dir)})" if e["output"] else name
            unit = UNITS.get(e["kind"])
            size = f"{len(e['blocks'])} {unit}s" if unit else human_size(e["size"])
            lines.append(f"| {link} | {KIND_LABEL[e['kind']]} | {size} | {e['words']:,} | {', '.join(e['flags'])} |")
            if e["outline"]:
                outlines += ["", f"### {src['title']}", ""] + [f"- {t}" for t in e["outline"][:30]]
            attention += [(src, f, e["note"]) for f in e["flags"] if f in FLAG_TEXT]
        if outlines:
            lines += ["", "## Outlines (slide titles / headings)"] + outlines
        write_if_changed(sec_dir / "_index.md", "\n".join(lines) + "\n")
        total = sum(state[s["key"]]["words"] for s in srcs)
        course_lines.append(f"| [{sec['section_number']:02d} {sec['section_name']}]"
                            f"({quote(srcs[0]['section_dir'])}/_index.md) | {len(srcs)} | {total:,} |")
    if attention:
        course_lines += ["", "## Needs attention", ""]
        course_lines += [f"- **{src['title']}** ({f}): {FLAG_TEXT[f]} {note}".rstrip() for src, f, note in attention]
    write_if_changed(text_root / "_index.md", "\n".join(course_lines) + "\n")


def build_passages(sources: list[dict], state: dict, manifest: dict) -> list[dict]:
    """Search units for Stage 5: one per page/slide; Word paragraphs grouped under their heading."""
    out = []
    for src in sources:
        e = state[src["key"]]
        base = {"source": src["key"], "md": e["output"], "section_number": src["section_number"],
                "section_name": src["section_name"], "title": src["title"], "kind": e["kind"]}
        if e["kind"] != "docx":
            out += [{**base, "location": b["loc"], "heading": b["heading"], "text": b["text"]}
                    for b in e["blocks"] if b["text"].strip()]
            continue
        group: list[dict] = []

        def flush():
            if group:
                loc = group[0]["loc"] if len(group) == 1 else f"{group[0]['loc']}-{group[-1]['loc']}"
                out.append({**base, "location": loc, "heading": group[-1]["heading"],
                            "text": "\n".join(b["text"] for b in group)})
                group.clear()

        for b in e["blocks"]:
            if b["level"] or sum(len(g["text"]) for g in group) > PASSAGE_CHARS:
                flush()
            group.append(b)
        flush()
    for item in manifest["items"]:
        if item["type"] == "url" and item.get("target_url"):
            out.append({"source": item["local_path"], "md": "", "section_number": item["section_number"],
                        "section_name": item["section_name"], "title": item["title"], "kind": "link",
                        "location": "link", "heading": "", "text": f"{item['title']}\n{item['target_url']}"})
    return out


def human_size(n: int) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024
    return ""


# --- index.html --------------------------------------------------------------------

def esc(s: str) -> str:
    return html.escape(s or "", quote=True)


def rich(md: str) -> str:
    """Escape, then turn **bold** into <strong>."""
    return re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", esc(md))


BADGES = {"needs_ocr": "needs OCR", "mostly_images": "mostly images", "needs_conversion": "not converted",
          "unsupported": "unsupported", "error": "not readable", "no_text": "no text"}

HTML_CSS = """
:root{--bg:#fbfaf7;--fg:#1d1d1b;--muted:#6b6a65;--line:#e4e1d8;--card:#fff;--accent:#2f5d8a;--warn:#9a5b00;
--warnbg:#fff3df;--mark:#ffe58a;--grey:#9b9a94}
@media (prefers-color-scheme:dark){:root{--bg:#141413;--fg:#ecebe6;--muted:#a19f97;--line:#2e2d2a;--card:#1c1c1a;
--accent:#8ab4e0;--warn:#f0b357;--warnbg:#33270f;--mark:#6b5a12;--grey:#77766f}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--fg);font:16px/1.55 system-ui,-apple-system,
"Segoe UI",sans-serif}main{max-width:920px;margin:0 auto;padding:24px 16px 80px;overflow-wrap:anywhere}
h1{font-size:1.6rem;margin:0 0 4px}h2{font-size:1.15rem;margin:36px 0 12px;padding-bottom:6px;
border-bottom:1px solid var(--line)}h3{font-size:1rem;margin:0}h4{font-size:.95rem;margin:18px 0 4px}
.sub{color:var(--muted);margin:0 0 16px}a{color:var(--accent)}
.bar{position:sticky;top:0;background:var(--bg);padding:10px 0;z-index:2;display:flex;gap:12px;align-items:center}
#q{flex:1;min-width:0;font:inherit;padding:8px 12px;border:1px solid var(--line);border-radius:8px;
background:var(--card);color:var(--fg)}#n{color:var(--muted);font-size:.9rem;white-space:nowrap}
.file{background:var(--card);border:1px solid var(--line);border-radius:10px;padding:12px 14px;margin:10px 0}
.file header{display:flex;flex-wrap:wrap;gap:8px;align-items:baseline}
.kind{font-size:.75rem;text-transform:uppercase;letter-spacing:.04em;color:var(--muted)}
.meta{color:var(--muted);font-size:.88rem;margin:4px 0 0}.meta a{margin-right:2px}
.badge{font-size:.75rem;padding:1px 8px;border-radius:99px;background:var(--warnbg);color:var(--warn)}
.badge.info{background:transparent;border:1px solid var(--line);color:var(--muted)}
.note{font-size:.88rem;color:var(--warn);margin:6px 0 0}
details{margin-top:8px}summary{cursor:pointer;color:var(--accent);font-size:.9rem}
.body{margin-top:8px;border-top:1px solid var(--line);padding-top:4px;overflow-wrap:anywhere}
.t{white-space:pre-wrap;margin:6px 0}.loc{color:var(--muted);font-size:.8rem}
.missing{opacity:.6}.missing h3{color:var(--grey)}
.links{padding-left:20px}.links li{margin:4px 0}
.attn{border:1px solid var(--line);border-radius:10px;padding:10px 14px;background:var(--card)}
.attn li{margin:4px 0}::highlight(m){background:var(--mark);color:inherit}
"""

HTML_JS = """
const q=document.getElementById('q'),n=document.getElementById('n');
const files=[...document.querySelectorAll('.file,.links li')],secs=[...document.querySelectorAll('.sec')];
const hl=(window.CSS&&CSS.highlights)?CSS.highlights:null;let timer;
q.addEventListener('input',()=>{clearTimeout(timer);timer=setTimeout(go,150)});
function go(){const term=q.value.trim().toLowerCase();let shown=0;const ranges=[];
files.forEach(f=>{const hit=!term||f.textContent.toLowerCase().includes(term);f.hidden=!hit;
if(!term||!hit)return;shown++;const d=f.querySelector('details');
if(d&&term.length>1)d.open=d.querySelector('.body').textContent.toLowerCase().includes(term);
if(hl&&term.length>1&&ranges.length<3000){const w=document.createTreeWalker(f,NodeFilter.SHOW_TEXT);let t;
while((t=w.nextNode())){const s=t.data.toLowerCase();let i=s.indexOf(term);
while(i>-1&&ranges.length<3000){const r=new Range();r.setStart(t,i);r.setEnd(t,i+term.length);ranges.push(r);
i=s.indexOf(term,i+term.length)}}}});
secs.forEach(s=>s.hidden=!!term&&!s.querySelector('.file:not([hidden]),.links li:not([hidden])'));
if(hl){hl.delete('m');if(ranges.length)hl.set('m',new Highlight(...ranges))}
n.textContent=term?shown+' matching':''}
"""


def render_body(entry: dict) -> str:
    kind, parts = entry["kind"], []
    for b in entry["blocks"]:
        if kind == "pdf":
            parts.append(f"<h4>Page {b['n']}</h4><div class=t>{rich(b['md'])}</div>")
        elif kind in ("pptx", "ppt"):
            parts.append(f"<h4>Slide {b['n']}: {esc(b['heading'] or '(untitled)')}</h4><div class=t>{rich(b['md'])}</div>")
        elif kind == "docx" and b["level"]:
            parts.append(f"<h4>{esc(b['md'])} <span class=loc>¶{b['n']}</span></h4>")
        elif kind == "docx":
            parts.append(f"<div class=t>{rich(b['md'])} <span class=loc>¶{b['n']}</span></div>")
        else:
            parts.append(f"<div class=t>{rich(b['md'])}</div>")
    return "".join(parts)


def build_html(manifest: dict, cdir: Path, sources: list[dict], state: dict) -> str:
    by_item: dict[str, list[dict]] = {}
    for src in sources:
        by_item.setdefault(src["item"]["id"], []).append(src)
    items_by_section: dict[int, list[dict]] = {}
    for it in manifest["items"]:
        items_by_section.setdefault(it["section_number"], []).append(it)

    total_words = sum(state[s["key"]]["words"] for s in sources)
    attention = []
    sections_html = []
    for sec in manifest["sections"]:
        cards, links = [], []
        for it in items_by_section.get(sec["section_number"], []):
            moodle = f' · <a href="{esc(it["url"])}" target=_blank rel=noopener>Moodle</a>' if it["url"] else ""
            if it["type"] == "url":
                if it.get("target_url"):
                    links.append(f'<li><a href="{esc(it["target_url"])}" target=_blank rel=noopener>{esc(it["title"])}</a>'
                                 f' <span class=loc>{esc(it["target_url"])}</span></li>')
                else:
                    links.append(f'<li class=missing>{esc(it["title"])} <span class=loc>({esc(it["status"])})</span></li>')
                continue
            srcs = by_item.get(it["id"], [])
            if not srcs:
                why = {"skipped": "Not downloaded (nothing to download)", "unavailable": "Restricted on Moodle",
                       "failed": "Download failed", "pending": "Not downloaded yet"}.get(it["status"], it["status"])
                note = f"<p class=note>{esc(it['note'])}</p>" if it["note"] and it["status"] != "skipped" else ""
                cards.append(f'<article class="file missing"><header><span class=kind>{esc(it["module"])}</span>'
                             f'<h3>{esc(it["title"])}</h3></header><p class=meta>{esc(why)}{moodle}</p>{note}</article>')
                continue
            for src in srcs:
                e = state[src["key"]]
                anchor = "f" + hashlib.sha1(src["key"].encode()).hexdigest()[:10]
                badges = "".join(f'<span class="badge{" info" if f in INFO_FLAGS else ""}">{BADGES[f]}</span>'
                                 for f in e["flags"])
                if e.get("converted"):
                    badges += '<span class="badge info">converted from .ppt</span>'
                unit = UNITS.get(e["kind"])
                bits = [f"{len(e['blocks'])} {unit}s" if unit else "", f"{e['words']:,} words" if e["words"] else "",
                        human_size(e["size"])]
                meta = " · ".join(b for b in bits if b)
                links_html = f' · <a href="{quote(src["rel_course"])}">Open file</a>'
                if e["output"]:
                    links_html += f' · <a href="{md_link(ROOT / e["output"], cdir)}">Text (.md)</a>'
                note = f"<p class=note>{esc(e['note'])}</p>" if e["note"] and set(e["flags"]) & set(FLAG_TEXT) else ""
                body = render_body(e)
                details = f"<details><summary>Show text</summary><div class=body>{body}</div></details>" if body else ""
                cards.append(f'<article class=file id={anchor}><header><span class=kind>{esc(KIND_LABEL[e["kind"]])}</span>'
                             f'<h3>{esc(src["title"])}</h3>{badges}</header>'
                             f'<p class=meta>{meta}{links_html}{moodle}</p>{note}{details}</article>')
                attention += [(anchor, src["title"], f) for f in e["flags"] if f in FLAG_TEXT]
        if not cards and not links:
            continue
        links_html = f"<h4>Links</h4><ul class=links>{''.join(links)}</ul>" if links else ""
        sections_html.append(f'<section class=sec id=s{sec["section_number"]}>'
                             f'<h2>{sec["section_number"]:02d} · {esc(sec["section_name"])}</h2>'
                             f'{"".join(cards)}{links_html}</section>')

    attn_html = ""
    if attention:
        rows = "".join(f'<li><a href="#{a}">{esc(t)}</a>: {esc(FLAG_TEXT[f])}</li>' for a, t, f in attention)
        attn_html = f"<div class=attn><strong>Needs attention</strong><ul>{rows}</ul></div>"
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
    return (f'<!doctype html><html lang=en><meta charset=utf-8>'
            f'<meta name=viewport content="width=device-width,initial-scale=1">'
            f'<title>{esc(manifest["course_name"])}</title><style>{HTML_CSS}</style><body><main>'
            f'<h1>{esc(manifest["course_name"])}</h1>'
            f'<p class=sub>{len(sources)} files · {total_words:,} words · extracted {stamp} · '
            f'generated by extract.py, offline</p>{attn_html}'
            f'<div class=bar><input id=q type=search placeholder="Filter files and text…" '
            f'aria-label="Filter"><span id=n></span></div>{"".join(sections_html)}'
            f'</main><script>{HTML_JS}</script></body></html>\n')


# --- main ---------------------------------------------------------------------------

def load_state(path: Path) -> dict:
    try:
        state = json.loads(path.read_text(encoding="utf-8"))
    except (FileNotFoundError, json.JSONDecodeError):
        return {"version": EXTRACTOR_VERSION, "sources": {}}
    if state.get("version") != EXTRACTOR_VERSION:
        log.info("Extractor changed since the last run: re-extracting everything.")
        return {"version": EXTRACTOR_VERSION, "sources": {}}
    return state


def is_current(src: dict, entry: dict | None) -> bool:
    """Same file version as last time, and its .md (if it has one) still exists."""
    if not entry or set(entry["flags"]) & RETRY_FLAGS:
        return False
    if entry["output"] and not (ROOT / entry["output"]).exists():
        return False
    st = src["path"].stat()
    if entry["size"] != st.st_size:
        return False
    if entry["mtime_ns"] != st.st_mtime_ns:
        if entry["sha1"] != sha1(src["path"]):
            return False
        entry["mtime_ns"] = st.st_mtime_ns  # touched but identical
    return True


def main(args) -> int:
    if not MANIFEST_FILE.exists():
        raise ConfigError("No data/manifest.json yet. Run:  python scrape.py")
    try:
        manifest = json.loads(MANIFEST_FILE.read_text(encoding="utf-8"))
    except json.JSONDecodeError:
        raise ConfigError("data/manifest.json is unreadable. Rerun:  python scrape.py")
    cdir = course_dir(manifest)
    if not cdir.exists():
        raise ConfigError(f"{rel(cdir)} doesn't exist yet. Run:  python scrape.py")
    text_root = TEXT_DIR / cdir.name
    converted_root = CONVERTED_DIR / cdir.name
    state_path = text_root / "extract_state.json"
    state = load_state(state_path)
    entries = state["sources"]

    sources, missing = collect_sources(manifest, cdir)
    log.info("Course: %s (%d downloaded files)", manifest["course_name"], len(sources))
    counts = {"extracted": 0, "unchanged": 0, "removed": 0}
    try:
        todo = [s for s in sources if args.force or not is_current(s, entries.get(s["key"]))]
        for i, src in enumerate(todo, 1):
            kind = kind_of(src["path"])
            log.info("[%d/%d] %s", i, len(todo), src["rel_course"])
            result = extract_source(src, kind, converted_root)
            st = src["path"].stat()
            out = output_path(text_root, src, kind) if result["blocks"] or kind in ("md", "txt", "html") else None
            entries[src["key"]] = {
                "kind": kind, "size": st.st_size, "mtime_ns": st.st_mtime_ns, "sha1": sha1(src["path"]),
                "output": rel(out) if out else "", "converted": result["converted"], "flags": result["flags"],
                "note": result["note"], "outline": result["outline"], "words": words(result["blocks"]),
                "blocks": result["blocks"],
            }
            for f in result["flags"]:
                if f in FLAG_TEXT:
                    log.warning("  %s: %s %s", f, FLAG_TEXT[f], result["note"])
            counts["extracted"] += 1
        counts["unchanged"] = len(sources) - len(todo)

        # Sources that are gone from the manifest or disk: remove what we generated for them.
        current = {s["key"] for s in sources}
        for key in [k for k in entries if k not in current]:
            for p in (entries[key]["output"], entries[key].get("converted")):
                if p:
                    (ROOT / p).unlink(missing_ok=True)
            del entries[key]
            counts["removed"] += 1
    finally:
        write_json_atomic(state_path, state)

    for src in sources:  # (re)write .md files whose content changed (incl. new titles after a recrawl)
        e = entries[src["key"]]
        if e["output"]:
            write_if_changed(ROOT / e["output"], render_md(src, e))
    write_indexes(manifest, text_root, sources, entries)
    passages = build_passages(sources, entries, manifest)
    write_text(text_root / "passages.jsonl", "".join(json.dumps(p, ensure_ascii=False) + "\n" for p in passages))
    index_html = cdir / "index.html"
    write_text(index_html, build_html(manifest, cdir, sources, entries))

    print_summary(manifest, sources, entries, counts, missing, text_root, len(passages), index_html)
    return 0


def print_summary(manifest, sources, entries, counts, missing, text_root, n_passages, index_html) -> None:
    kinds: dict[str, int] = {}
    for s in sources:
        k = entries[s["key"]]["kind"]
        kinds[k] = kinds.get(k, 0) + 1
    print()
    print(f"Course: {manifest['course_name']}")
    print(f"Files: {len(sources)}   extracted now {counts['extracted']}, unchanged {counts['unchanged']}, "
          f"removed {counts['removed']}")
    print("By type: " + ", ".join(f"{KIND_LABEL[k]} {n}" for k, n in sorted(kinds.items(), key=lambda x: -x[1])))
    print(f"Words: {sum(entries[s['key']]['words'] for s in sources):,}")
    converted = [s for s in sources if entries[s["key"]].get("converted")]
    if converted:
        print("Converted with PowerPoint: " + ", ".join(s["path"].name for s in converted))
    media = [s for s in sources if "no_text" in entries[s["key"]]["flags"]]
    if media:
        print("No text (media): " + ", ".join(s["path"].name for s in media))
    attention = [(s, f) for s in sources for f in entries[s["key"]]["flags"] if f in FLAG_TEXT]
    if attention:
        print("Needs attention:")
        for s, f in attention:
            print(f"  - {s['path'].name} [{f}]: {FLAG_TEXT[f]} {entries[s['key']]['note']}".rstrip())
    if missing:
        print(f"Missing on disk ({len(missing)}), rerun python scrape.py to fetch them again:")
        for p in missing:
            print(f"  - {p}")
    print()
    print(f"Text:    {(text_root / '_index.md').relative_to(ROOT)}")
    print(f"Search:  {(text_root / 'passages.jsonl').relative_to(ROOT)}  ({n_passages} passages)")
    print(f"Open:    {index_html.relative_to(ROOT)}")


if __name__ == "__main__":
    parser = base_parser("Extract clean text from the downloaded course files and build index.html.")
    parser.add_argument("--force", action="store_true", help="re-extract every file, even unchanged ones")
    run(main, parser.parse_args())
