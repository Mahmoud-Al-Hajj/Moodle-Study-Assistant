"""Stage 4 tests: extract.py on sample documents generated at run time (no binary fixtures).

Run with the rest:  python tests\\run_tests.py
Or alone:           python tests\\test_extract.py
"""

import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
import unittest
from pathlib import Path

PROJECT = Path(__file__).resolve().parents[1]
SCRIPTS = ["common.py", "moodle_selectors.py", "downloader.py", "extract.py"]
COURSE = "Test Course"


def png_bytes() -> bytes:
    import pymupdf
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 40), False)
    pix.clear_with(180)
    return pix.tobytes("png")


def make_pptx(path: Path, image_slides: int = 0) -> None:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])  # title + content
    s.shapes.title.text = "Capitalism basics"
    tf = s.placeholders[1].text_frame
    tf.text = "Invisible hand"
    p = tf.add_paragraph()
    p.text = "Self-interest benefits society"
    p.level = 1
    group = s.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(5), Inches(3), Inches(1)).text_frame.text = "Grouped note text"
    table = s.shapes.add_table(2, 2, Inches(5), Inches(5), Inches(3), Inches(1)).table
    for (r, c), v in {(0, 0): "Term", (0, 1): "Meaning", (1, 0): "Market", (1, 1): "Exchange place"}.items():
        table.cell(r, c).text = v
    s.notes_slide.notes_text_frame.text = "Remember the pin factory example."
    for _ in range(image_slides):
        blank = prs.slides.add_slide(prs.slide_layouts[6])
        blank.shapes.add_picture(io.BytesIO(png_bytes()), Inches(1), Inches(1))
    prs.save(str(path))


def make_docx(path: Path) -> None:
    import docx
    d = docx.Document()
    d.add_heading("Will to Power", level=1)
    p = d.add_paragraph()
    p.add_run("The will to power").bold = True
    p.add_run(" is defined as the drive to grow.")
    d.add_paragraph("First list point", style="List Bullet")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Concept", "Meaning"
    t.cell(1, 0).text, t.cell(1, 1).text = "Übermensch", "Higher human"
    d.add_heading("Arabic", level=2)
    d.add_paragraph("الفرز والبحث")
    d.save(str(path))


def make_pdf(path: Path, scanned: bool = False) -> None:
    import pymupdf
    doc = pymupdf.open()
    for i in range(2):
        page = doc.new_page()
        if scanned:
            page.insert_image(pymupdf.Rect(50, 50, 300, 300), stream=png_bytes())
        else:
            page.insert_text((72, 72), f"Human rights page {i + 1}: dignity and freedom for every person.")
    doc.save(str(path))


class TestExtract(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp(prefix="moodle-extract-"))
        for name in SCRIPTS:
            shutil.copy(PROJECT / name, self.dir / name)
        self.course = self.dir / "downloads" / COURSE
        s1, s2 = self.course / "01-Week 1", self.course / "02-Week 2"
        s1.mkdir(parents=True)
        s2.mkdir(parents=True)
        make_pptx(s1 / "slides.pptx")
        make_pptx(s1 / "photos.pptx", image_slides=3)
        make_docx(s1 / "guide.docx")
        make_pdf(s2 / "rights.pdf")
        make_pdf(s2 / "scan.pdf", scanned=True)
        (s2 / "old.ppt").write_bytes(b"not really a ppt")
        (s2 / "lecture.mp4").write_bytes(b"\x00" * 2048)
        (s2 / "Essay - instructions.md").write_text("# Essay\n\nSource: x\n\nWrite 500 words.\n", encoding="utf-8")
        (self.course / "links.md").write_text("# Links\n", encoding="utf-8")

        def item(i, sec, title, module, paths, **extra):
            local = paths[0] if len(paths) == 1 else str(Path(paths[0]).parent.as_posix())
            return {"section_number": sec, "section_name": f"Week {sec}", "title": title,
                    "type": module if module in ("resource", "folder", "page", "url", "assign") else "other",
                    "url": f"https://moodle.example/mod/{module}/view.php?id={i}", "local_path": local,
                    "status": "downloaded", "id": f"cm:{i}", "module": module, "note": "",
                    "files": paths if len(paths) > 1 or module == "assign" else [], "target_url": "", **extra}

        d = f"downloads/{COURSE}"
        self.items = [
            item(1, 1, "Slides <script>alert(1)</script>", "resource", [f"{d}/01-Week 1/slides.pptx"]),
            item(2, 1, "Photo deck", "resource", [f"{d}/01-Week 1/photos.pptx"]),
            item(3, 1, "Nietzsche guide", "resource", [f"{d}/01-Week 1/guide.docx"]),
            item(4, 2, "Rights", "resource", [f"{d}/02-Week 2/rights.pdf"]),
            item(5, 2, "Scanned reading", "resource", [f"{d}/02-Week 2/scan.pdf"]),
            item(6, 2, "Old deck", "resource", [f"{d}/02-Week 2/old.ppt"]),
            item(7, 2, "Lecture video", "resource", [f"{d}/02-Week 2/lecture.mp4"]),
            item(8, 2, "Essay", "assign", [f"{d}/02-Week 2/Essay - instructions.md"]),
            item(9, 2, "Visualgo", "url", [f"{d}/links.md"], target_url="https://example.org/v"),
            {**item(10, 2, "Broken file", "resource", [""]), "status": "failed", "local_path": "",
             "note": "HTTP 404 for x"},
        ]
        self.write_manifest()

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def write_manifest(self):
        manifest = {"moodle_url": "https://moodle.example", "course_id": "42", "course_name": COURSE,
                    "crawled_at": "2026-10-17T00:00:00+00:00",
                    "sections": [{"section_number": 1, "section_name": "Week 1", "item_count": 3},
                                 {"section_number": 2, "section_name": "Week 2", "item_count": 7}],
                    "items": self.items}
        (self.dir / "data").mkdir(exist_ok=True)
        (self.dir / "data" / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")

    def run_extract(self, *args) -> tuple[int, str]:
        # PowerShell "missing": .ppt conversion must fail gracefully (real PowerPoint isn't used in tests).
        env = {**os.environ, "PYTHONIOENCODING": "utf-8", "MOODLE_STUDY_POWERSHELL": "no-such-powershell-exe"}
        proc = subprocess.run([sys.executable, "extract.py", *args], cwd=self.dir, env=env,
                              capture_output=True, text=True, encoding="utf-8", timeout=180)
        out = proc.stdout + proc.stderr
        self.assertNotIn("Traceback", out)
        return proc.returncode, out

    def text(self, rel: str) -> str:
        return (self.dir / "data" / "text" / COURSE / rel).read_text(encoding="utf-8")

    def state(self) -> dict:
        return json.loads(self.text("extract_state.json"))["sources"]

    def entry(self, name: str) -> dict:
        return next(e for k, e in self.state().items() if k.endswith(name))

    def test_full_run(self):
        code, out = self.run_extract()
        self.assertEqual(code, 0, out)
        self.assertIn("extracted now 8", out)

        slides = self.text("01-Week 1/slides.pptx.md")
        self.assertIn("## Slide 1: Capitalism basics", slides)
        self.assertIn("- Invisible hand", slides)
        self.assertIn("  - Self-interest benefits society", slides)  # indented level 1
        self.assertIn("Grouped note text", slides)
        self.assertIn("| Market | Exchange place |", slides)
        self.assertIn("### Notes", slides)
        self.assertIn("pin factory", slides)
        self.assertEqual(self.entry("photos.pptx")["flags"], ["mostly_images"])
        self.assertIn("No text on this slide", self.text("01-Week 1/photos.pptx.md"))

        guide = self.text("01-Week 1/guide.docx.md")
        self.assertIn("## Will to Power [¶1]", guide)
        self.assertIn("**The will to power** is defined as the drive to grow. [¶2]", guide)
        self.assertIn("- First list point [¶3]", guide)
        self.assertIn("[¶4]\n\n| Concept | Meaning |", guide)
        self.assertIn("| Übermensch | Higher human |", guide)
        self.assertIn("الفرز والبحث", guide)

        rights = self.text("02-Week 2/rights.pdf.md")
        self.assertIn("## Page 2", rights)
        self.assertIn("dignity and freedom", rights)
        self.assertEqual(self.entry("scan.pdf")["flags"], ["needs_ocr"])
        self.assertEqual(self.entry("old.ppt")["flags"], ["needs_conversion"])
        self.assertEqual(self.entry("lecture.mp4")["flags"], ["no_text"])
        self.assertIn("Write 500 words.", self.text("02-Week 2/Essay - instructions.md"))
        self.assertNotIn("old.ppt.md", str(list((self.dir / "data" / "text").rglob("*"))))

        self.assertIn("Needs attention", self.text("_index.md"))
        self.assertIn("Capitalism basics", self.text("01-Week 1/_index.md"))  # outline
        for f in ("needs_ocr", "mostly_images", "needs_conversion"):
            self.assertIn(f, out)

    def test_passages(self):
        self.run_extract()
        lines = [json.loads(x) for x in self.text("passages.jsonl").splitlines()]
        for p in lines:
            self.assertTrue({"source", "location", "text", "title", "section_number", "kind"} <= set(p))
        locs = {(Path(p["source"]).name, p["location"]) for p in lines}
        self.assertIn(("slides.pptx", "slide 1"), locs)
        self.assertIn(("rights.pdf", "page 2"), locs)
        self.assertIn(("links.md", "link"), locs)
        docx = [p for p in lines if p["source"].endswith("guide.docx")]
        self.assertEqual([p["location"] for p in docx], ["¶1-¶4", "¶5-¶6"])  # grouped under headings
        self.assertEqual(docx[1]["heading"], "Arabic")

    def test_rerun_and_changes(self):
        self.run_extract()
        md = self.dir / "data" / "text" / COURSE / "01-Week 1" / "slides.pptx.md"
        mtime = md.stat().st_mtime_ns
        code, out = self.run_extract()
        self.assertIn("extracted now 1, unchanged 7", out)  # only old.ppt retried (needs_conversion)
        self.assertEqual(md.stat().st_mtime_ns, mtime)

        time.sleep(0.05)
        guide = self.course / "01-Week 1" / "guide.docx"
        os.utime(guide)  # touched, same content: not re-extracted
        code, out = self.run_extract()
        self.assertIn("extracted now 1, unchanged 7", out)

        make_pdf(self.course / "02-Week 2" / "rights.pdf", scanned=True)  # content changed
        code, out = self.run_extract()
        self.assertIn("extracted now 2, unchanged 6", out)
        self.assertEqual(self.entry("rights.pdf")["flags"], ["needs_ocr"])

        self.items = [it for it in self.items if it["id"] != "cm:4"]  # gone from the manifest
        self.write_manifest()
        code, out = self.run_extract()
        self.assertIn("removed 1", out)
        self.assertFalse((self.dir / "data" / "text" / COURSE / "02-Week 2" / "rights.pdf.md").exists())

    def test_retitled_item_rewrites_md_without_reextracting(self):
        self.run_extract()
        self.items[0]["title"] = "Renamed slides"
        self.write_manifest()
        code, out = self.run_extract()
        self.assertIn("unchanged 7", out)
        self.assertIn("# Renamed slides", self.text("01-Week 1/slides.pptx.md"))

    def test_index_html(self):
        self.run_extract()
        page = (self.course / "index.html").read_text(encoding="utf-8")
        for s in ("Week 1", "Week 2", "Photo deck", "Nietzsche guide", "Lecture video", "Visualgo",
                  "https://example.org/v", "Broken file", "HTTP 404 for x", "needs OCR", "Capitalism basics"):
            self.assertIn(s, page)
        self.assertNotIn("<script>alert(1)</script>", page)
        self.assertIn("Slides &lt;script&gt;alert(1)&lt;/script&gt;", page)
        self.assertNotIn("<script src", page)
        self.assertNotIn('<link href="http', page)
        self.assertIn('href="01-Week%201/slides.pptx"', page)
        self.assertEqual(page.count("<script>"), 1)  # only the inline filter script

    def test_missing_manifest(self):
        (self.dir / "data" / "manifest.json").unlink()
        code, out = self.run_extract()
        self.assertEqual(code, 1)
        self.assertIn("python scrape.py", out)


if __name__ == "__main__":
    unittest.main(verbosity=2)
